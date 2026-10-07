import os
import re
import openpyxl
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls

BASE_DIR = '/Users/tomwagener/development/WSF Office'
EXCEL_PATH = os.path.join(BASE_DIR, 'Member_Database.xlsx')
TEMPLATE_PATH = os.path.join(BASE_DIR, 'Invoice Template.docx')
OUTPUT_DIR = os.path.join(BASE_DIR, 'Invoices_2027')
OVERVIEW_EXCEL_PATH = os.path.join(BASE_DIR, 'Invoices_2027_Overview.xlsx')

os.makedirs(OUTPUT_DIR, exist_ok=True)

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb['Organizations']

def sanitize_filename(name):
    s = re.sub(r'[^\w\s-]', '', name).strip()
    s = re.sub(r'[-\s]+', '_', s)
    return s

def format_address(addr, nation):
    lines = []
    if addr and addr.strip() not in ['-', 'N/A']:
        parts = [p.strip() for p in addr.split(',') if p.strip()]
        if len(parts) <= 3:
            lines = parts
        elif len(parts) == 4:
            lines = [f'{parts[0]}, {parts[1]}', parts[2], parts[3]]
        elif len(parts) == 5:
            lines = [f'{parts[0]}, {parts[1]}', f'{parts[2]}, {parts[3]}', parts[4]]
        else:
            lines = [f'{parts[0]}, {parts[1]}', ', '.join(parts[2:-1]), parts[-1]]
            
        if nation and nation.strip() not in ['-', 'N/A']:
            nat = nation.strip()
            last_line_lower = lines[-1].lower() if lines else ''
            if nat.lower() not in last_line_lower:
                lines.append(nat)
    else:
        if nation and nation.strip() not in ['-', 'N/A']:
            lines.append(nation.strip())
            
    return lines

members = []
for r in range(2, ws.max_row + 1):
    org_id = ws.cell(row=r, column=1).value
    name = ws.cell(row=r, column=2).value
    addr = ws.cell(row=r, column=3).value
    vat_id = ws.cell(row=r, column=4).value
    status = ws.cell(row=r, column=5).value
    gen_email = ws.cell(row=r, column=6).value
    bill_email = ws.cell(row=r, column=8).value
    mem_type = ws.cell(row=r, column=17).value
    fee = ws.cell(row=r, column=18).value
    nation = ws.cell(row=r, column=19).value
    
    if fee is not None and isinstance(fee, (int, float)) and status != 'Candiate' and org_id != 'WSF037' and name:
        members.append({
            'org_id': str(org_id).strip(),
            'name': str(name).strip(),
            'addr': str(addr).strip() if addr else '',
            'vat_id': str(vat_id).strip() if vat_id else '',
            'email': str(bill_email or gen_email or '').strip(),
            'mem_type': str(mem_type).strip() if mem_type else '',
            'fee': float(fee),
            'nation': str(nation).strip() if nation else '',
            'status': str(status).strip()
        })

print(f'Found {len(members)} eligible members for 2027 invoices.')

summary_rows = []

for m in members:
    org_num_match = re.search(r'\d+', m['org_id'])
    org_num = int(org_num_match.group(0)) if org_num_match else 0
    invoice_number = f'0127{org_num:02d}'
    
    invoice_date = 'January 01, 2027'
    payment_due = 'March 02, 2027'
    amount_eur = f'€ {m["fee"]:,.2f}'
    
    doc = docx.Document(TEMPLATE_PATH)
    body = doc.element.body
    
    # 1. Replace P3-P8 with a 2-column borderless table
    p3_elem = doc.paragraphs[3]._p
    idx = body.index(p3_elem)
    
    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    
    tblBorders = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    tbl._tbl.tblPr.append(tblBorders)
    
    row = tbl.rows[0]
    row.cells[0].width = Inches(3.8)
    row.cells[1].width = Inches(2.45)
    
    # Left Cell: BILL TO
    c_left = row.cells[0]
    addr_lines = format_address(m['addr'], m['nation'])
    
    left_entries = [('BILL TO', False)]
    left_entries.append((m['name'], True))
    for line in addr_lines:
        left_entries.append((line, False))
    if m['vat_id'] and m['vat_id'] not in ['-', 'N/A', 'None']:
        left_entries.append((f'VAT: {m["vat_id"]}', False))
    if m['email']:
        left_entries.append((m['email'], False))
        
    for i, (txt, is_bold) in enumerate(left_entries):
        p = c_left.paragraphs[0] if i == 0 else c_left.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r_run = p.add_run(txt)
        r_run.bold = is_bold
        r_run.font.name = 'Arial'
        r_run.font.size = Pt(11)
        r_run.font.color.rgb = RGBColor(0, 0, 0)
        
    # Right Cell: Invoice Details
    c_right = row.cells[1]
    right_entries = [
        ('Invoice Number: ', invoice_number),
        ('Invoice Date: ', invoice_date),
        ('Payment Due: ', payment_due),
        ('Amount Due (EUR): ', amount_eur)
    ]
    for i, (lbl, val) in enumerate(right_entries):
        p = c_right.paragraphs[0] if i == 0 else c_right.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(lbl)
        r1.bold = True
        r1.font.name = 'Arial'
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(0, 0, 0)
        r2 = p.add_run(val)
        r2.bold = False
        r2.font.name = 'Arial'
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(0, 0, 0)
        
    body.insert(idx, tbl._tbl)
    
    # Remove original template paragraphs P3 to P8 (6 paragraphs)
    for _ in range(6):
        body.remove(doc.paragraphs[3]._p)
        
    # 2. Update Items Table (now doc.tables[1])
    items_tbl = doc.tables[1]
    widths = [4500, 1200, 1650, 1679]  # 9029 total
    grid = items_tbl._tbl.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblGrid')
    if grid is not None:
        for i, col in enumerate(grid.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridCol')):
            col.set(qn('w:w'), str(widths[i]))
            
    for r_idx in [0, 1]:
        for c_idx, cell in enumerate(items_tbl.rows[r_idx].cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcW')
            if tcW is not None:
                tcW.set(qn('w:w'), str(widths[c_idx]))
                
    items_tbl.rows[1].cells[0].paragraphs[0].text = 'WSF Membership 2027'
    items_tbl.rows[1].cells[1].paragraphs[0].text = '1'
    items_tbl.rows[1].cells[2].paragraphs[0].text = amount_eur
    items_tbl.rows[1].cells[3].paragraphs[0].text = amount_eur
    
    # Remove second placeholder item row
    if len(items_tbl.rows) > 2:
        items_tbl._tbl.remove(items_tbl.rows[2]._tr)
        
    # 3. Update TOTAL line
    for p in doc.paragraphs:
        if 'TOTAL' in p.text:
            p.text = ''
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            r_tot = p.add_run(f'\t\t\t\t\t\tTOTAL\t\t{amount_eur}')
            r_tot.bold = True
            r_tot.font.name = 'Arial'
            r_tot.font.size = Pt(12)
            r_tot.underline = True
            
    # 4. Clean up empty paragraphs between TOTAL and Notes / Terms
    total_found = False
    empty_to_keep = 3
    empty_kept = 0
    for p in list(doc.paragraphs):
        if 'TOTAL' in p.text:
            total_found = True
            continue
        if total_found:
            if 'Notes / Terms' in p.text:
                break
            if p.text.strip() == '':
                if empty_kept < empty_to_keep:
                    if p._p.pPr is not None:
                        rPr = p._p.pPr.find(qn('w:rPr'))
                        if rPr is not None:
                            u = rPr.find(qn('w:u'))
                            if u is not None:
                                rPr.remove(u)
                    empty_kept += 1
                else:
                    body.remove(p._p)
                    
    # 5. Footer: dark gray
    for s in doc.sections:
        for p in s.footer.paragraphs:
            for r_run in p.runs:
                r_run.font.color.rgb = RGBColor(80, 80, 80)
                
    filename = f'Invoice_{invoice_number}_{sanitize_filename(m["name"])}.docx'
    out_path = os.path.join(OUTPUT_DIR, filename)
    doc.save(out_path)
    
    summary_rows.append({
        'Amount Euro': m['fee'],
        'Customer Name': m['name'],
        'Invoice Number': invoice_number,
        'Type': 'Membership',
        'Org ID': m['org_id'],
        'Filename': filename
    })

print(f'Successfully generated {len(summary_rows)} invoices in {OUTPUT_DIR}.')

# Create summary Excel file
sum_wb = openpyxl.Workbook()
sum_ws = sum_wb.active
sum_ws.title = 'Invoices 2027'

headers = ['Amount Euro', 'Customer Name', 'Invoice Number', 'Type']
sum_ws.append(headers)

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

header_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')
header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
data_font = Font(name='Arial', size=10)
bold_font = Font(name='Arial', size=10, bold=True)

thin_border = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

for col_idx, col_name in enumerate(headers, 1):
    cell = sum_ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

total_amount = 0.0
for r_idx, row_data in enumerate(summary_rows, 2):
    c1 = sum_ws.cell(row=r_idx, column=1, value=row_data['Amount Euro'])
    c1.number_format = '€#,##0.00'
    c1.font = data_font
    c1.alignment = Alignment(horizontal='right')
    c1.border = thin_border
    
    c2 = sum_ws.cell(row=r_idx, column=2, value=row_data['Customer Name'])
    c2.font = data_font
    c2.alignment = Alignment(horizontal='left')
    c2.border = thin_border
    
    c3 = sum_ws.cell(row=r_idx, column=3, value=row_data['Invoice Number'])
    c3.font = data_font
    c3.alignment = Alignment(horizontal='center')
    c3.border = thin_border
    
    c4 = sum_ws.cell(row=r_idx, column=4, value=row_data['Type'])
    c4.font = data_font
    c4.alignment = Alignment(horizontal='center')
    c4.border = thin_border
    
    total_amount += row_data['Amount Euro']

# Total row
tot_row = len(summary_rows) + 2
tot_c1 = sum_ws.cell(row=tot_row, column=1, value=total_amount)
tot_c1.number_format = '€#,##0.00'
tot_c1.font = bold_font
tot_c1.alignment = Alignment(horizontal='right')
tot_c1.border = thin_border

tot_c2 = sum_ws.cell(row=tot_row, column=2, value='Total')
tot_c2.font = bold_font
tot_c2.border = thin_border

tot_c3 = sum_ws.cell(row=tot_row, column=3, value=f'{len(summary_rows)} Invoices')
tot_c3.font = bold_font
tot_c3.alignment = Alignment(horizontal='center')
tot_c3.border = thin_border

tot_c4 = sum_ws.cell(row=tot_row, column=4, value='')
tot_c4.border = thin_border

for col in sum_ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    sum_ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

sum_wb.save(OVERVIEW_EXCEL_PATH)
print(f'Saved overview Excel to {OVERVIEW_EXCEL_PATH}. Total: €{total_amount:,.2f}')
